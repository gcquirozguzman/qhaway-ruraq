"""Extracción de texto de presentaciones, documentos y transcripciones a Markdown."""

from __future__ import annotations

import re
from pathlib import Path


def leer_texto(ruta: Path) -> str:
    datos = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return datos.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", errors="replace")


# --- pptx --------------------------------------------------------------------
def _textos_shape(shape) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    salida: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            salida += _textos_shape(s)
    elif getattr(shape, "has_table", False) and shape.has_table:
        for fila in shape.table.rows:
            celdas = [c.text.strip().replace("\n", " ") for c in fila.cells]
            salida.append("| " + " | ".join(celdas) + " |")
    elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            texto = "".join(r.text for r in p.runs).strip()
            if texto:
                salida.append("  " * p.level + "- " + texto)
    return salida


def extraer_pptx(ruta: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(ruta))
    partes = [f"# {ruta.name}\n"]
    for i, slide in enumerate(prs.slides, 1):
        titulo = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            titulo = slide.shapes.title.text_frame.text.strip()
        partes.append(f"## Diapositiva {i}" + (f": {titulo}" if titulo else ""))
        lineas = []
        for shape in slide.shapes:
            if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
                continue
            lineas += _textos_shape(shape)
        if lineas:
            partes.append("\n".join(lineas))
        if slide.has_notes_slide:
            notas = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notas:
                partes.append(f"**Notas del orador:** {notas}")
        partes.append("")
    return "\n\n".join(partes)


# --- pdf ---------------------------------------------------------------------
def extraer_pdf(ruta: Path) -> str:
    import pymupdf

    partes = [f"# {ruta.name}\n"]
    with pymupdf.open(str(ruta)) as doc:
        for i, pagina in enumerate(doc, 1):
            texto = pagina.get_text("text").strip()
            if texto:
                partes.append(f"## Página {i}\n\n{texto}")
    if len(partes) == 1:
        partes.append("_(PDF sin texto extraíble: probablemente escaneado)_")
    return "\n\n".join(partes)


# --- docx --------------------------------------------------------------------
def extraer_docx(ruta: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(ruta))
    partes = [f"# {ruta.name}\n"]
    for bloque in d.element.body.iterchildren():
        tag = bloque.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(bloque, d)
            texto = p.text.strip()
            if not texto:
                continue
            estilo = (p.style.name or "").lower() if p.style is not None else ""
            m = re.search(r"(heading|título|titulo)\s*(\d)", estilo)
            if m:
                partes.append("#" * min(int(m.group(2)) + 1, 6) + " " + texto)
            elif "list" in estilo or "lista" in estilo:
                partes.append("- " + texto)
            else:
                partes.append(texto)
        elif tag == "tbl":
            t = Table(bloque, d)
            for fila in t.rows:
                celdas = [c.text.strip().replace("\n", " ") for c in fila.cells]
                partes.append("| " + " | ".join(celdas) + " |")
    return "\n\n".join(partes)


# --- transcripciones ---------------------------------------------------------
RE_TIEMPO = re.compile(r"(\d{1,2}:)?\d{2}:\d{2}[.,]\d{3}\s*-->")


def extraer_vtt(ruta: Path) -> str:
    """Compacta VTT/SRT a líneas '[hh:mm:ss] texto' para ahorrar tokens."""
    lineas = [f"# Transcripción: {ruta.name}\n"]
    tiempo = ""
    previo = ""
    for linea in leer_texto(ruta).splitlines():
        linea = linea.strip()
        if not linea or linea == "WEBVTT" or linea.isdigit() or linea.startswith("NOTE"):
            continue
        if RE_TIEMPO.search(linea):
            inicio = linea.split("-->")[0].strip().replace(",", ".")
            tiempo = inicio.split(".")[0]
            if tiempo.count(":") == 1:
                tiempo = "00:" + tiempo
            continue
        texto = re.sub(r"<[^>]+>", "", linea)
        if texto and texto != previo:
            lineas.append(f"[{tiempo}] {texto}")
            previo = texto
    return "\n".join(lineas)


EXTRACTORES = {
    ".pptx": extraer_pptx,
    ".pdf": extraer_pdf,
    ".docx": extraer_docx,
    ".vtt": extraer_vtt,
    ".srt": extraer_vtt,
    ".txt": lambda r: f"# {r.name}\n\n" + leer_texto(r),
    ".md": leer_texto,
}


def extraer(ruta: Path) -> str:
    funcion = EXTRACTORES.get(ruta.suffix.lower())
    if not funcion:
        raise ValueError(f"Tipo no soportado: {ruta.suffix}")
    return funcion(ruta).strip() + "\n"
